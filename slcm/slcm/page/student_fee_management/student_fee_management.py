import frappe
from frappe import _
from frappe.utils import cint, flt, today

ALLOWED_ROLES = ("System Manager", "Campus Admin", "Accounts Manager", "Accounts User")

# Demand-level aggregate per student. Cancelled demands never count toward dues.
_DEMAND_AGG = """
	SELECT
		student,
		COUNT(*)                                                         AS demand_count,
		SUM(net_payable)                                                 AS total_payable,
		SUM(paid_amount)                                                 AS paid_amount,
		SUM(waiver_amount)                                               AS waiver_amount,
		SUM(outstanding_amount)                                          AS outstanding_amount,
		SUM(CASE WHEN status = 'Overdue' THEN outstanding_amount ELSE 0 END) AS overdue_amount
	FROM `tabFee Demand`
	WHERE status != 'Cancelled' {component_condition}
	GROUP BY student
"""

# Excess = unused balance on active (submitted) credit notes.
_CREDIT_AGG = """
	SELECT student, SUM(available_credit) AS excess_amount
	FROM `tabStudent Credit Note`
	WHERE docstatus = 1 AND status = 'Active'
	GROUP BY student
"""


def _check_access():
	frappe.only_for(ALLOWED_ROLES)


def _require_remarks(remarks):
	remarks = (remarks or "").strip()
	if not remarks:
		frappe.throw(_("Remarks are mandatory."))
	return remarks


def _as_list(value):
	"""Accept a list, a JSON-encoded list, or a single plain value; drop blanks."""
	if not value:
		return []
	if isinstance(value, str):
		value = frappe.parse_json(value) if value.lstrip().startswith("[") else [value]
	return [v for v in value if v not in (None, "")]


@frappe.whitelist()
def get_filter_options():
	"""Academic years, terms (with their year, for dependent filtering) and programmes present on students."""
	_check_access()

	years = frappe.db.sql_list(
		"""SELECT DISTINCT academic_year FROM `tabStudent Master`
		WHERE IFNULL(academic_year, '') != '' ORDER BY academic_year DESC"""
	)
	# Terms set up in Academic Term plus any term held on a student (students copy theirs from the batch,
	# so a batch without a term leaves students blank — the master list keeps the filter usable).
	terms = frappe.db.sql(
		"""SELECT academic_year, name AS academic_term FROM `tabAcademic Term`
		WHERE IFNULL(status, '') != 'Inactive'
		UNION
		SELECT DISTINCT academic_year, academic_term FROM `tabStudent Master`
		WHERE IFNULL(academic_term, '') != ''
		ORDER BY academic_term""",
		as_dict=True,
	)
	programmes = frappe.db.sql(
		"""SELECT DISTINCT sm.programme_of_study AS name, p.program_name
		FROM `tabStudent Master` sm
		LEFT JOIN `tabProgramme` p ON p.name = sm.programme_of_study
		WHERE IFNULL(sm.programme_of_study, '') != ''
		ORDER BY sm.programme_of_study""",
		as_dict=True,
	)
	fee_components = frappe.db.sql_list(
		"""SELECT DISTINCT fee_component FROM `tabFee Demand`
		WHERE status != 'Cancelled' AND IFNULL(fee_component, '') != '' ORDER BY fee_component"""
	)
	return {"academic_years": years, "terms": terms, "programmes": programmes, "fee_components": fee_components}


@frappe.whitelist()
def get_students(
	academic_year=None,
	academic_term=None,
	programme=None,
	dues_status=None,
	fee_component=None,
	search=None,
	start=0,
	page_length=50,
	sort_by=None,
	sort_order=None,
):
	"""Paginated student list with per-student fee totals, plus totals across the whole filtered set."""
	_check_access()

	base, values = _student_base(academic_year, academic_term, programme, dues_status, fee_component, search)

	totals = frappe.db.sql(
		f"""SELECT
			COUNT(*)                              AS student_count,
			SUM(IFNULL(fd.total_payable, 0))      AS total_payable,
			SUM(IFNULL(fd.paid_amount, 0))        AS paid_amount,
			SUM(IFNULL(fd.outstanding_amount, 0)) AS outstanding_amount,
			SUM(IFNULL(fd.overdue_amount, 0))     AS overdue_amount,
			SUM(IFNULL(cn.excess_amount, 0))      AS excess_amount
		{base}""",
		values,
		as_dict=True,
	)[0]

	values.update(start=cint(start), page_length=min(cint(page_length) or 50, 500))
	rows = frappe.db.sql(
		f"""SELECT
			sm.name AS student, sm.first_name AS student_name, sm.registration_id,
			sm.official_email_id, sm.email, sm.programme_of_study, sm.batch,
			sm.academic_year, sm.academic_term, sm.section, sm.student_status,
			IFNULL(fd.demand_count, 0)       AS demand_count,
			IFNULL(fd.total_payable, 0)      AS total_payable,
			IFNULL(fd.paid_amount, 0)        AS paid_amount,
			IFNULL(fd.waiver_amount, 0)      AS waiver_amount,
			IFNULL(fd.outstanding_amount, 0) AS outstanding_amount,
			IFNULL(fd.overdue_amount, 0)     AS overdue_amount,
			IFNULL(cn.excess_amount, 0)      AS excess_amount
		{base}
		ORDER BY {_order_by(sort_by, sort_order)}
		LIMIT %(start)s, %(page_length)s""",
		values,
		as_dict=True,
	)
	_attach_receipt_info(rows)

	return {"rows": rows, "totals": totals}


def _student_base(academic_year=None, academic_term=None, programme=None, dues_status=None, fee_component=None, search=None):
	"""FROM/WHERE clause (aliases sm, fd, cn) selecting the students the list filters match,
	plus its bind values. Shared by the student list and the reports so they always agree."""
	conditions = ["1=1"]
	values = {}
	# Each filter takes one or more values (multi-select); empty means "no filter".
	for param, column, key in (
		(academic_year, "sm.academic_year", "academic_year"),
		(academic_term, "sm.academic_term", "academic_term"),
		(programme, "sm.programme_of_study", "programme"),
	):
		selected = _as_list(param)
		if selected:
			conditions.append(f"{column} IN %({key})s")
			values[key] = tuple(selected)
	if search:
		conditions.append(
			"(sm.name LIKE %(search)s OR sm.first_name LIKE %(search)s "
			"OR sm.registration_id LIKE %(search)s OR sm.official_email_id LIKE %(search)s)"
		)
		values["search"] = f"%{search.strip()}%"

	status_conditions = {
		"pending": "IFNULL(fd.outstanding_amount, 0) > 0",
		"overdue": "IFNULL(fd.overdue_amount, 0) > 0",
		"cleared": "IFNULL(fd.demand_count, 0) > 0 AND IFNULL(fd.outstanding_amount, 0) = 0",
		"excess": "IFNULL(cn.excess_amount, 0) > 0",
		"no_demands": "IFNULL(fd.demand_count, 0) = 0",
	}
	# Several statuses → a student matching any of them is included.
	chosen = [status_conditions[k] for k in _as_list(dues_status) if k in status_conditions]
	if chosen:
		conditions.append("(" + " OR ".join(f"({c})" for c in chosen) + ")")

	# Fee Component: only students with dues for these components, and every amount counts only them
	components = _as_list(fee_component)
	component_condition = ""
	if components:
		component_condition = "AND fee_component IN %(fee_component)s"
		values["fee_component"] = tuple(components)
		conditions.append("IFNULL(fd.demand_count, 0) > 0")
	demand_agg = _DEMAND_AGG.format(component_condition=component_condition)

	base = f"""
		FROM `tabStudent Master` sm
		LEFT JOIN ({demand_agg}) fd ON fd.student = sm.name
		LEFT JOIN ({_CREDIT_AGG}) cn ON cn.student = sm.name
		WHERE {" AND ".join(conditions)}
	"""
	return base, values


# Sortable columns → SQL expressions. Only these keys are accepted, so sort input never reaches SQL raw.
_SORT_COLUMNS = {
	"student": ["sm.first_name"],
	"programme": ["sm.programme_of_study"],
	"year_term": ["sm.academic_year", "sm.academic_term"],
	"batch": ["sm.batch"],
	"demand_count": ["IFNULL(fd.demand_count, 0)"],
	"total_payable": ["IFNULL(fd.total_payable, 0)"],
	"paid_amount": ["IFNULL(fd.paid_amount, 0)"],
	"outstanding_amount": ["IFNULL(fd.outstanding_amount, 0)"],
	"excess_amount": ["IFNULL(cn.excess_amount, 0)"],
	# Same precedence as the list's status badge: Overdue > Pending > Cleared > No Demands
	"status": [
		"""CASE WHEN IFNULL(fd.demand_count, 0) = 0 THEN 0
			WHEN IFNULL(fd.overdue_amount, 0) > 0 THEN 3
			WHEN IFNULL(fd.outstanding_amount, 0) > 0 THEN 2
			ELSE 1 END"""
	],
}


def _order_by(sort_by, sort_order):
	if sort_by not in _SORT_COLUMNS:
		return "IFNULL(fd.outstanding_amount, 0) DESC, sm.first_name ASC, sm.name ASC"
	direction = "ASC" if (sort_order or "").lower() == "asc" else "DESC"
	return ", ".join(f"{col} {direction}" for col in _SORT_COLUMNS[sort_by]) + ", sm.name ASC"


def _attach_receipt_info(rows):
	"""Add receipt_count and latest_receipt (active Fee Receipts) to each student row on the page."""
	if not rows:
		return
	receipts = frappe.get_all(
		"Fee Receipt",
		filters={"student": ["in", [r.student for r in rows]], "status": ["!=", "Cancelled"]},
		fields=["name", "student"],
		order_by="receipt_date desc, creation desc",
	)
	by_student = {}
	for rc in receipts:
		by_student.setdefault(rc.student, []).append(rc.name)
	for r in rows:
		names = by_student.get(r.student, [])
		r["receipt_count"] = len(names)
		r["latest_receipt"] = names[0] if names else None


@frappe.whitelist()
def get_student_receipts(student):
	"""Active Fee Receipts for one student, newest first, with the demands each one settled."""
	_check_access()

	receipts = frappe.get_all(
		"Fee Receipt",
		filters={"student": student, "status": ["!=", "Cancelled"]},
		fields=["name", "receipt_date", "amount", "payment_mode", "reference_number", "fee_payment"],
		order_by="receipt_date desc, creation desc",
	)
	if receipts:
		paid = frappe.get_all(
			"Fee Receipt Demands Paid",
			filters={"parent": ["in", [r.name for r in receipts]], "parenttype": "Fee Receipt"},
			fields=["parent", "fee_demand", "description"],
			order_by="idx asc",
		)
		by_parent = {}
		for row in paid:
			by_parent.setdefault(row.parent, []).append(row.description or row.fee_demand)
		for r in receipts:
			r["demands"] = by_parent.get(r.name, [])
	return receipts


@frappe.whitelist()
def download_receipt(receipt):
	"""
	PDF for a Fee Receipt, rendered exactly as staff print it today: the linked Fee Payment in the
	"Fee Payment Receipt - Admin Copy" format (the Fee Payment form's Print Receipt → Admin Copy).
	Receipts with no linked Fee Payment fall back to the student portal's orphan-receipt layout.
	Access is gated by this page's roles, since Accounts roles lack read on Fee Receipt / Fee Payment.
	"""
	_check_access()

	row = frappe.db.get_value("Fee Receipt", receipt, ["status", "fee_payment"], as_dict=True)
	if not row:
		frappe.throw(_("Receipt {0} not found.").format(receipt))
	if row.status == "Cancelled":
		frappe.throw(_("Receipt {0} has been cancelled.").format(receipt))

	fee_payment = row.fee_payment or frappe.db.get_value("Fee Payment", {"receipt": receipt}, "name")
	if fee_payment:
		frappe.flags.ignore_print_permissions = True
		try:
			pdf = frappe.get_print(
				"Fee Payment", fee_payment, "Fee Payment Receipt - Admin Copy", as_pdf=True, no_letterhead=0
			)
		finally:
			frappe.flags.ignore_print_permissions = False
	else:
		from slcm.api.student_portal import _generate_orphan_receipt_pdf

		pdf = _generate_orphan_receipt_pdf(receipt)

	frappe.local.response.filename = "{0}.pdf".format(receipt.replace(" ", "-").replace("/", "-"))
	frappe.local.response.filecontent = pdf
	frappe.local.response.type = "pdf"


@frappe.whitelist()
def get_student_dues(student):
	"""Everything the student detail view needs: profile, summary, demands, payments, credits, refunds."""
	_check_access()

	if not frappe.db.exists("Student Master", student):
		frappe.throw(_("Student {0} not found.").format(student))

	profile = frappe.db.get_value(
		"Student Master",
		student,
		[
			"name", "first_name", "registration_id", "official_email_id", "email", "phone",
			"programme_of_study", "batch", "academic_year", "academic_term", "section",
			"student_status", "passport_size_photo",
			"applying_scholarship", "scholarship_type", "scholarship_amount", "scholarship_percentage",
			"scholarship_approval_date", "fee_waiver_remarks",
		],
		as_dict=True,
	)
	if profile.programme_of_study:
		profile.update(
			frappe.db.get_value(
				"Programme", profile.programme_of_study, ["program_name", "department"], as_dict=True
			)
			or {}
		)

	demands = frappe.get_all(
		"Fee Demand",
		filters={"student": student},
		fields=[
			"name", "demand_type", "fee_component", "description", "status",
			"demand_date", "due_date", "creation", "academic_year",
			"original_amount", "penalty_amount", "waiver_amount", "net_payable",
			"paid_amount", "credit_adjusted", "outstanding_amount", "refunded_amount", "moved_to_excess_amount",
			"is_refundable", "last_payment_date", "remarks",
		],
		order_by="demand_date desc, creation desc",
	)

	# Receipts that settled each demand (a receipt can cover several demands, a demand several receipts)
	receipt_rows = frappe.db.sql(
		"""SELECT rdp.fee_demand, fr.name, fr.receipt_date, fr.payment_mode, rdp.amount
		FROM `tabFee Receipt Demands Paid` rdp
		JOIN `tabFee Receipt` fr ON fr.name = rdp.parent AND rdp.parenttype = 'Fee Receipt'
		WHERE fr.student = %s AND IFNULL(fr.status, '') != 'Cancelled'
		ORDER BY fr.receipt_date DESC, fr.creation DESC""",
		student,
		as_dict=True,
	)
	by_demand = {}
	for r in receipt_rows:
		by_demand.setdefault(r.pop("fee_demand"), []).append(r)
	for d in demands:
		d["receipts"] = by_demand.get(d.name, [])

	payments = frappe.get_all(
		"Fee Payment",
		filters={"student": student},
		fields=[
			"name", "payment_date", "payment_mode", "amount", "reference_number",
			"transaction_date", "status", "docstatus", "receipt", "remarks", "university_bank_account",
		],
		order_by="payment_date desc, creation desc",
	)
	if payments:
		rows = frappe.get_all(
			"Fee Payment Demand Row",
			filters={"parent": ["in", [p.name for p in payments]], "parenttype": "Fee Payment"},
			fields=["parent", "fee_demand", "demand_description", "amount_allocated"],
			order_by="idx asc",
		)
		by_parent = {}
		for r in rows:
			by_parent.setdefault(r.parent, []).append(r)
		for p in payments:
			p["allocations"] = by_parent.get(p.name, [])

	credit_notes = frappe.get_all(
		"Student Credit Note",
		filters={"student": student, "docstatus": 1},
		fields=[
			"name", "credit_type", "academic_year", "credit_amount", "available_credit",
			"used_credit", "status", "source_receipt", "remarks", "creation",
		],
		order_by="creation desc",
	)
	if credit_notes:
		adjustments = frappe.get_all(
			"Credit Adjustment Row",
			filters={"parent": ["in", [c.name for c in credit_notes]], "parenttype": "Student Credit Note"},
			fields=["parent", "fee_demand", "fee_component", "amount_adjusted", "adjusted_on", "adjusted_by"],
			order_by="idx asc",
		)
		by_parent = {}
		for a in adjustments:
			by_parent.setdefault(a.parent, []).append(a)
		for c in credit_notes:
			c["adjustments"] = by_parent.get(c.name, [])

	refunds = frappe.get_all(
		"Fee Refund",
		filters={"student": student},
		fields=[
			"name", "fee_demand", "fee_component", "refund_type", "refund_amount",
			"refund_date", "refund_mode", "utr_number", "status", "docstatus", "reason",
		],
		order_by="refund_date desc, creation desc",
	)

	concessions = frappe.get_all(
		"Fee Concession",
		filters={"student": student, "docstatus": ["!=", 2]},
		fields=[
			"name", "fee_demand", "fee_component", "concession_type",
			"waiver_value", "waiver_amount", "status", "docstatus", "reason", "approved_on", "creation",
		],
		order_by="creation desc",
	)
	stipends = frappe.get_all(
		"Stipend Payment",
		filters={"student": student, "docstatus": ["!=", 2]},
		fields=[
			"name", "stipend_type", "academic_year", "academic_term", "payment_date", "payment_mode",
			"amount", "reference_number", "status", "docstatus", "remarks",
		],
		order_by="payment_date desc, creation desc",
	)

	active = [d for d in demands if d.status != "Cancelled"]
	summary = {
		"total_payable": sum(flt(d.net_payable) for d in active),
		"paid_amount": sum(flt(d.paid_amount) for d in active),
		"waiver_amount": sum(flt(d.waiver_amount) for d in active),
		"penalty_amount": sum(flt(d.penalty_amount) for d in active),
		"pending_amount": sum(flt(d.outstanding_amount) for d in active),
		"overdue_amount": sum(flt(d.outstanding_amount) for d in active if d.status == "Overdue"),
		"excess_amount": sum(
			flt(c.available_credit) for c in credit_notes if c.status == "Active"
		),
		"refunded_amount": sum(flt(r.refund_amount) for r in refunds if r.docstatus == 1),
		"scholarship_amount": sum(flt(c.waiver_amount) for c in concessions if c.status == "Approved"),
		"stipend_paid": sum(flt(x.amount) for x in stipends if x.docstatus == 1),
	}

	return {
		"profile": profile,
		"summary": summary,
		"demands": demands,
		"payments": payments,
		"credit_notes": credit_notes,
		"refunds": refunds,
		"concessions": concessions,
		"stipends": stipends,
	}


@frappe.whitelist()
def record_payment(
	student,
	allocations,
	payment_mode,
	payment_date=None,
	reference_number=None,
	transaction_date=None,
	bank_name=None,
	remarks=None,
	university_bank_account=None,
	settlement_date=None,
):
	"""
	Record one Fee Payment against one or more of a student's demands and submit it.
	allocations: JSON list of {"fee_demand": ..., "amount": ...}.
	Submitting runs the existing Fee Payment pipeline (demand status update, receipt, payment log).
	"""
	_check_access()
	remarks = _require_remarks(remarks)

	allocations = frappe.parse_json(allocations) or []
	allocations = [a for a in allocations if flt(a.get("amount")) > 0]
	if not allocations:
		frappe.throw(_("Enter an amount for at least one demand."))

	demand_names = [a["fee_demand"] for a in allocations]
	demands = {
		d.name: d
		for d in frappe.get_all(
			"Fee Demand",
			filters={"name": ["in", demand_names]},
			fields=["name", "student", "description", "fee_component", "outstanding_amount", "status"],
		)
	}
	for a in allocations:
		d = demands.get(a["fee_demand"])
		if not d or d.student != student:
			frappe.throw(_("Fee Demand {0} does not belong to this student.").format(a["fee_demand"]))
		if d.status in ("Paid", "Cancelled", "Waived"):
			frappe.throw(_("Fee Demand {0} is {1} and cannot take a payment.").format(d.name, d.status))

	payment = frappe.new_doc("Fee Payment")
	payment.student = student
	payment.payment_date = payment_date or today()
	payment.payment_mode = payment_mode
	payment.reference_number = reference_number
	payment.transaction_date = transaction_date
	payment.bank_name = bank_name
	payment.remarks = remarks
	payment.university_bank_account = university_bank_account
	payment.settlement_date = settlement_date or None
	payment.amount = sum(flt(a["amount"]) for a in allocations)
	for a in allocations:
		d = demands[a["fee_demand"]]
		payment.append("payment_demands", {
			"fee_demand": d.name,
			"demand_description": d.description or d.fee_component,
			"outstanding_amount": d.outstanding_amount,
			"amount_allocated": flt(a["amount"]),
		})
	payment.insert()
	payment.submit()

	return {"payment": payment.name, "receipt": payment.receipt}


@frappe.whitelist()
def apply_excess_credit(student, fee_demand, amount, remarks=None):
	"""Adjust a student's excess (active credit notes, oldest first) against one demand."""
	_check_access()
	remarks = _require_remarks(remarks)

	amount = flt(amount)
	if amount <= 0:
		frappe.throw(_("Amount must be greater than zero."))

	demand = frappe.db.get_value(
		"Fee Demand", fee_demand, ["student", "outstanding_amount", "status"], as_dict=True
	)
	if not demand or demand.student != student:
		frappe.throw(_("Fee Demand {0} does not belong to this student.").format(fee_demand))
	if demand.status in ("Paid", "Cancelled", "Waived"):
		frappe.throw(_("Fee Demand {0} is {1}.").format(fee_demand, demand.status))
	if amount > flt(demand.outstanding_amount):
		frappe.throw(
			_("Amount ({0}) exceeds the outstanding amount ({1}).").format(
				frappe.utils.fmt_money(amount, currency="INR"),
				frappe.utils.fmt_money(demand.outstanding_amount, currency="INR"),
			)
		)

	notes = frappe.get_all(
		"Student Credit Note",
		filters={"student": student, "docstatus": 1, "status": "Active", "available_credit": [">", 0]},
		fields=["name", "available_credit"],
		order_by="creation asc",
	)
	if amount > sum(flt(n.available_credit) for n in notes):
		frappe.throw(_("Amount exceeds the student's available excess."))

	remaining = amount
	used = []
	for n in notes:
		if remaining <= 0:
			break
		take = min(remaining, flt(n.available_credit))
		frappe.get_doc("Student Credit Note", n.name).apply_credit_to_demand(fee_demand, take)
		used.append(n.name)
		remaining -= take

	frappe.get_doc("Fee Demand", fee_demand).add_comment(
		"Comment",
		_("Adjusted {0} from excess ({1}). Remarks: {2}").format(
			frappe.utils.fmt_money(amount, currency="INR"), ", ".join(used), frappe.utils.escape_html(remarks)
		),
	)
	return {"credit_notes": used}


@frappe.whitelist()
def cancel_demand(fee_demand, remarks=None):
	_check_access()
	remarks = _require_remarks(remarks)
	demand = frappe.get_doc("Fee Demand", fee_demand)
	result = demand.cancel_demand()
	demand.add_comment("Comment", _("Due cancelled. Remarks: {0}").format(frappe.utils.escape_html(remarks)))
	return result


@frappe.whitelist()
def move_to_excess(student, fee_demand, amount, credit_type, remarks=None, academic_year=None, cancel_due=0):
	"""
	Move money already paid against a due into the student's excess.

	The amount comes off the due's paid amount (the due's outstanding/status recalculate through the
	standard Fee Demand logic, same as a refund) and a submitted Student Credit Note is created for it,
	so it shows under Excess and can later be adjusted against other dues or refunded.
	Optionally cancels the due when its whole paid amount has been moved.
	"""
	_check_access()
	remarks = _require_remarks(remarks)
	amount = flt(amount)
	if amount <= 0:
		frappe.throw(_("Amount must be greater than zero."))

	demand = frappe.get_doc("Fee Demand", fee_demand)
	if demand.student != student:
		frappe.throw(_("Fee Demand {0} does not belong to this student.").format(fee_demand))
	if demand.status == "Cancelled":
		frappe.throw(_("Fee Demand {0} is cancelled.").format(fee_demand))
	if amount > flt(demand.paid_amount):
		frappe.throw(
			_("Amount ({0}) exceeds the amount paid on {1} ({2}).").format(
				frappe.utils.fmt_money(amount, currency="INR"),
				fee_demand,
				frappe.utils.fmt_money(demand.paid_amount, currency="INR"),
			)
		)
	if cint(cancel_due) and amount < flt(demand.paid_amount):
		frappe.throw(_("A due can only be cancelled when its whole paid amount is moved to excess."))

	source_receipt = frappe.db.sql(
		"""SELECT fr.name FROM `tabFee Receipt Demands Paid` rdp
		JOIN `tabFee Receipt` fr ON fr.name = rdp.parent AND rdp.parenttype = 'Fee Receipt'
		WHERE rdp.fee_demand = %s AND IFNULL(fr.status, '') != 'Cancelled'
		ORDER BY fr.receipt_date DESC, fr.creation DESC LIMIT 1""",
		fee_demand,
	)

	demand.update_payment_status(paid_delta=-amount)

	note = frappe.new_doc("Student Credit Note")
	note.student = student
	note.credit_type = credit_type
	note.academic_year = academic_year or demand.academic_year
	note.credit_amount = amount
	note.source_receipt = source_receipt[0][0] if source_receipt else None
	note.remarks = _("Moved from due {0} ({1}). {2}").format(
		fee_demand, demand.fee_component or demand.description or "", remarks
	)
	note.insert()
	note.submit()

	# Status: "Moved to Excess", or Cancelled (shown as "Cancelled & Moved to Excess" via moved_to_excess_amount)
	demand.reload()
	moved_total = flt(demand.moved_to_excess_amount) + amount
	if cint(cancel_due):
		demand.cancel_demand()
		demand.db_set("moved_to_excess_amount", moved_total)
	else:
		demand.db_set({"moved_to_excess_amount": moved_total, "status": "Moved to Excess"})

	demand.add_comment(
		"Comment",
		_("{0} moved to excess as {1}. Remarks: {2}").format(
			frappe.utils.fmt_money(amount, currency="INR"), note.name, frappe.utils.escape_html(remarks)
		),
	)
	return {"credit_note": note.name}


@frappe.whitelist()
def record_stipend(
	student,
	stipend_type,
	payment_date,
	payment_mode,
	amount,
	remarks=None,
	academic_year=None,
	academic_term=None,
	reference_number=None,
):
	"""Create and submit a Stipend Payment for the student (normal Stipend Payment permissions apply)."""
	_check_access()
	remarks = _require_remarks(remarks)
	if flt(amount) <= 0:
		frappe.throw(_("Amount must be greater than zero."))

	sm = frappe.db.get_value("Student Master", student, ["first_name", "programme_of_study"], as_dict=True)
	if not sm:
		frappe.throw(_("Student {0} not found.").format(student))

	doc = frappe.new_doc("Stipend Payment")
	doc.update({
		"student": student,
		"student_name": sm.first_name,
		"programme": sm.programme_of_study if frappe.db.exists("Programme", sm.programme_of_study) else None,
		"academic_year": academic_year,
		"academic_term": academic_term,
		"stipend_type": stipend_type,
		"payment_date": payment_date,
		"payment_mode": payment_mode,
		"amount": flt(amount),
		"reference_number": reference_number,
		"remarks": remarks,
		"status": "Draft",
	})
	doc.insert()
	doc.submit()
	return {"stipend_payment": doc.name}


# ─────────────────────────────────────────────────────────────────────────────
# Reports (Due Wise / Student Wise Outstanding) — same filters as the student list.
# Column headers follow the office's sheets in ~/Desktop/SLCM.
# ─────────────────────────────────────────────────────────────────────────────

# (key, sheet header, is_amount)
DUE_WISE_COLUMNS = [
	("student_id", "Student ID", False),
	("student_name", "Student Name", False),
	("email", "Email", False),
	("programme", "Programme", False),
	("batch", "Batch", False),
	("due_type", "Due Type", False),
	("voucher_number", "Voucher Number", False),
	("fee_component", "Fee Component", False),
	("remarks", "Remarks", False),
	("original_amount", "Original Amount", True),
	("waiver_amount", "Scholarship / Waiver Amount", True),
	("penalty_amount", "Penalty Amount", True),
	("net_payable", "Net Payable", True),
	("paid_amount", "Amount Paid", True),
	("outstanding_amount", "Amount Outstanding", True),
	("due_status", "Due Status", False),
	("demand_date", "Demand Date", False),
	("due_date", "Due Date", False),
	("settlement_date", "Settlement Date", False),
]
_DUE_AMOUNT_KEYS = [key for key, _label, is_amount in DUE_WISE_COLUMNS if is_amount]


def _report_students_sql(filters):
	base, values = _student_base(
		filters.get("academic_year"),
		filters.get("academic_term"),
		filters.get("programme"),
		filters.get("dues_status"),
		None,
		filters.get("search"),
	)
	return f"SELECT sm.name {base}", values


def _due_wise(filters, start=None, page_length=None):
	"""One row per (non-cancelled) Fee Demand of the filtered students."""
	students_sql, values = _report_students_sql(filters)
	where = f"d.status != 'Cancelled' AND d.student IN ({students_sql})"
	totals = frappe.db.sql(
		f"""SELECT COUNT(*) AS row_count,
			{", ".join(f"SUM(IFNULL(d.{k}, 0)) AS {k}" for k in _DUE_AMOUNT_KEYS)}
		FROM `tabFee Demand` d WHERE {where}""",
		values,
		as_dict=True,
	)[0]
	limit = ""
	if page_length:
		values.update(start=cint(start), page_length=min(cint(page_length), 500))
		limit = "LIMIT %(start)s, %(page_length)s"
	rows = frappe.db.sql(
		f"""SELECT
			d.student,
			COALESCE(NULLIF(sm.registration_id, ''), sm.name) AS student_id,
			COALESCE(NULLIF(d.student_name, ''), sm.first_name) AS student_name,
			COALESCE(NULLIF(d.student_email, ''), NULLIF(sm.official_email_id, ''), sm.email) AS email,
			sm.programme_of_study AS programme, sm.batch,
			d.demand_type AS due_type, d.name AS voucher_number, d.fee_component, d.remarks,
			{", ".join(f"IFNULL(d.{k}, 0) AS {k}" for k in _DUE_AMOUNT_KEYS)},
			d.status AS due_status, d.demand_date, d.due_date,
			CASE WHEN IFNULL(d.outstanding_amount, 0) <= 0 AND IFNULL(d.paid_amount, 0) > 0
				THEN d.last_payment_date END AS settlement_date
		FROM `tabFee Demand` d
		JOIN `tabStudent Master` sm ON sm.name = d.student
		WHERE {where}
		ORDER BY sm.first_name, sm.name, d.demand_date, d.name
		{limit}""",
		values,
		as_dict=True,
	)
	return rows, totals


def _outstanding(filters, start=None, page_length=None):
	"""One row per filtered student: outstanding per fee component (every Fee Component
	is a column, as in the sheet) and the total."""
	students_sql, values = _report_students_sql(filters)
	components = frappe.get_all("Fee Component", pluck="name", order_by="name asc")

	count = frappe.db.sql(f"SELECT COUNT(*) FROM ({students_sql}) s", values)[0][0]
	limit = ""
	if page_length:
		values.update(start=cint(start), page_length=min(cint(page_length), 500))
		limit = "LIMIT %(start)s, %(page_length)s"
	students = frappe.db.sql(
		f"""SELECT sm.name AS student,
			COALESCE(NULLIF(sm.registration_id, ''), sm.name) AS student_id,
			sm.first_name AS student_name,
			COALESCE(NULLIF(sm.official_email_id, ''), sm.email) AS email,
			sm.academic_status
		FROM `tabStudent Master` sm
		WHERE sm.name IN ({students_sql})
		ORDER BY sm.first_name, sm.name
		{limit}""",
		values,
		as_dict=True,
	)

	def by_component(student_condition, params):
		return frappe.db.sql(
			f"""SELECT student, fee_component, SUM(IFNULL(outstanding_amount, 0)) AS amount
			FROM `tabFee Demand`
			WHERE status != 'Cancelled' AND student IN {student_condition}
			GROUP BY student, fee_component""",
			params,
			as_dict=True,
		)

	page_amounts = by_component("%(page_students)s", {"page_students": tuple(s.student for s in students) or ("",)})
	grid = {}
	for a in page_amounts:
		grid.setdefault(a.student, {})[a.fee_component] = flt(a.amount)
		if a.fee_component not in components:
			components.append(a.fee_component)

	for s in students:
		s["amounts"] = {c: grid.get(s.student, {}).get(c, 0) for c in components}
		s["total"] = sum(s["amounts"].values())

	# Totals across every filtered student, not just this page.
	totals = {c: 0 for c in components}
	for a in by_component(f"({students_sql})", values):
		totals[a.fee_component] = totals.get(a.fee_component, 0) + flt(a.amount)
	totals["total"] = sum(v for k, v in totals.items() if k != "total")
	return components, students, count, totals


def _report_filters(academic_year, academic_term, programme, dues_status, search):
	return {
		"academic_year": academic_year,
		"academic_term": academic_term,
		"programme": programme,
		"dues_status": dues_status,
		"search": (search or "").strip(),
	}


@frappe.whitelist()
def get_due_wise_report(
	academic_year=None, academic_term=None, programme=None, dues_status=None, search=None, start=0, page_length=25
):
	_check_access()
	filters = _report_filters(academic_year, academic_term, programme, dues_status, search)
	rows, totals = _due_wise(filters, start, page_length or 25)
	return {"columns": [[k, label, a] for k, label, a in DUE_WISE_COLUMNS], "rows": rows, "totals": totals}


@frappe.whitelist()
def get_outstanding_report(
	academic_year=None, academic_term=None, programme=None, dues_status=None, search=None, start=0, page_length=25
):
	_check_access()
	filters = _report_filters(academic_year, academic_term, programme, dues_status, search)
	components, rows, count, totals = _outstanding(filters, start, page_length or 25)
	return {"components": components, "rows": rows, "count": count, "totals": totals}


@frappe.whitelist()
def export_report(report, academic_year=None, academic_term=None, programme=None, dues_status=None, search=None):
	"""Download the whole (unpaginated) report as .xlsx in the sheet's layout."""
	_check_access()
	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.utils import get_column_letter

	filters = _report_filters(academic_year, academic_term, programme, dues_status, search)
	wb = Workbook()
	ws = wb.active
	bold = Font(bold=True)
	header_fill = PatternFill("solid", fgColor="F2F2F2")
	amount_format = "#,##,##0.00"

	if report == "due_wise":
		ws.title = "Due Wise Report"
		rows, totals = _due_wise(filters)
		ws.append([label for _k, label, _a in DUE_WISE_COLUMNS])
		for r in rows:
			ws.append([
				(flt(r[k]) if is_amount else (r[k] if r[k] is not None else "")) for k, _label, is_amount in DUE_WISE_COLUMNS
			])
		total_row = ["Total"] + [""] * (len(DUE_WISE_COLUMNS) - 1)
		for i, (k, _label, is_amount) in enumerate(DUE_WISE_COLUMNS):
			if is_amount:
				total_row[i] = flt(totals.get(k))
		ws.append(total_row)
		amount_cols = [i + 1 for i, c in enumerate(DUE_WISE_COLUMNS) if c[2]]
		date_cols = [i + 1 for i, c in enumerate(DUE_WISE_COLUMNS) if c[0].endswith("_date")]
		filename = f"Due Wise Report {today()}.xlsx"
	elif report == "outstanding":
		ws.title = "Student wise outstanding fee"
		components, rows, _count, totals = _outstanding(filters)
		ws.append(["Sl. No.", "Student ID", "Student Name", "Student Email ID", "Academic Status"] + components + ["Total Outstanding Fee"])
		for i, r in enumerate(rows, 1):
			ws.append(
				[i, r.student_id, r.student_name or "", r.email or "", r.academic_status or ""]
				+ [flt(r["amounts"][c]) for c in components]
				+ [flt(r["total"])]
			)
		ws.append(["Total", "", "", "", ""] + [flt(totals.get(c)) for c in components] + [flt(totals["total"])])
		amount_cols = list(range(6, 6 + len(components) + 1))
		date_cols = []
		filename = f"Student wise outstanding fee report {today()}.xlsx"
	else:
		frappe.throw(_("Unknown report"))

	for cell in ws[1]:
		cell.font = bold
		cell.fill = header_fill
		cell.alignment = Alignment(wrap_text=True, vertical="center")
	for cell in ws[ws.max_row]:
		cell.font = bold
	for col in amount_cols:
		for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
			row[0].number_format = amount_format
	for col in date_cols:
		for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
			row[0].number_format = "DD-MM-YYYY"
	for i in range(1, ws.max_column + 1):
		ws.column_dimensions[get_column_letter(i)].width = 18
	ws.freeze_panes = "A2"

	out = BytesIO()
	wb.save(out)
	frappe.local.response.filename = filename
	frappe.local.response.filecontent = out.getvalue()
	frappe.local.response.type = "binary"
