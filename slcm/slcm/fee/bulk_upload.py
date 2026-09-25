"""Bulk-upload Excel/CSV templates for Fee Concession, Fee Payment and Fee Demand.

Each template mirrors the office's own spreadsheet (Desktop/SLCM/*.xlsx):
  * Fee Concession / Fee Payment — one row per open due, every column pre-filled from the system except
    the yellow ones the office fills in.
  * Fee Demand — a blank creation sheet (Voucher No. is assigned by the system), with dropdowns and a
    "Students" reference sheet so the office can look up each Student ID.

Headers equal the doctype field labels, so a filled sheet imports through Data Import
(Insert New Records) without any column mapping. Data Import's own "Download Template" returns these
sheets for the three doctypes (see ``data_import_download_template`` and hooks.py).
"""

from io import BytesIO

import frappe
from frappe import _
from frappe.utils import flt, today

# Dues columns shared by the Concession and Payment sheets: (header, width, key in the dues row)
DUES_COLUMNS = [
	("Student Registration Id / Application Number", 38, "registration_id"),
	("Student Email Id", 30, "student_email"),
	("Fee Component", 39, "fee_component"),
	("Voucher Number", 19, "name"),
	("Date of creation", 17, "demand_date"),
	("Academic Year", 13, "academic_year"),
	("Due Date", 14, "due_date"),
	("Remark", 60, "remark"),
	("Original Amount", 13, "original_amount"),
	("Penalty Amount", 13, "penalty_amount"),
	("Scholarship / Waiver Amount", 14, "waiver_amount"),
	("Total Payable", 13, "net_payable"),
	("Paid Amount", 13, "paid_amount"),
	("Pending Amount", 14, "outstanding_amount"),
]

# Input columns (yellow) per pre-filled template: (header, width, kind)
#   kind: "money" | "date" | "select:<fieldname>" (dropdown from that field's options) | None (free text)
PREFILLED = {
	"Fee Concession": {
		"filename": "Concession Bulk Upload",
		"inputs": [
			("Concession Type", 30, "select:concession_type"),
			("Waiver Value", 17, "money"),
			("Date of concession", 17, "date"),
			("Reason", 40, None),
		],
		# a due can carry only one approved concession
		"extra_conditions": [
			"""NOT EXISTS (SELECT 1 FROM `tabFee Concession` fc
				WHERE fc.fee_demand = fd.name AND fc.status = 'Approved' AND fc.docstatus = 1)"""
		],
	},
	"Fee Payment": {
		"filename": "Fee Payment Bulk Upload",
		"inputs": [
			("Payment Amount", 16, "money"),
			("Payment Date", 15, "date"),
			("Settlement Date", 15, "date"),
			("Payment Mode", 18, "select:payment_mode"),
			("Reference Number", 24, None),
			("University bank account", 38, "select:university_bank_account"),
			("Remarks", 40, None),
		],
		"extra_conditions": [],
	},
}

# Fee Demand creation sheet: (header, width, required, kind)
DEMAND_COLUMNS = [
	("Voucher No.", 16, False, "auto"),
	("Student ID", 18, True, "student"),
	("Student Email ID", 30, True, None),
	("Student Name", 26, False, None),
	("Academic Year", 14, False, "academic_year"),
	("Fee Component", 32, True, "fee_component"),
	("Original Amount", 14, True, "money"),
	("Waiver Amount", 14, True, "money"),
	("Penalty Amount", 14, True, "money"),
	("Net Payable", 14, True, "net"),
	("Demand Date", 14, True, "date"),
	("Due Date", 14, True, "date"),
	("Status", 12, False, None),
	("Description", 40, True, None),
]
DEMAND_ROWS = 500  # formatted / validated rows in the blank sheet

OPEN_DUES_CONDITIONS = [
	"fd.status NOT IN ('Paid', 'Cancelled', 'Waived')",
	"IFNULL(fd.outstanding_amount, 0) > 0",
]


def as_list(value):
	"""A multi-select filter value: list, JSON list or single value; empty means "all"."""
	if not value:
		return []
	if isinstance(value, str):
		value = frappe.parse_json(value) if value.lstrip().startswith("[") else [value]
	return [v for v in value if v]


def _check(doctype):
	if doctype not in PREFILLED and doctype != "Fee Demand":
		frappe.throw(_("No bulk-upload template for {0}.").format(doctype))
	if not frappe.has_permission(doctype, "create"):
		frappe.throw(_("You are not permitted to create {0}.").format(_(doctype)), frappe.PermissionError)


def _dues_conditions(doctype):
	return OPEN_DUES_CONDITIONS + PREFILLED[doctype]["extra_conditions"]


# ── Filter options ───────────────────────────────────────────────────────────
@frappe.whitelist()
def get_filter_options(doctype):
	"""Values for the template filters, with counts. Pre-filled sheets: years / programmes / fee components
	that have open dues. Fee Demand sheet: years / programmes of students (for the Students reference list)."""
	_check(doctype)
	out = {}
	if doctype == "Fee Demand":
		for key, col in (("academic_year", "academic_year"), ("programme", "programme_of_study")):
			out[key] = frappe.db.sql(
				f"""SELECT {col} AS value, COUNT(*) AS n FROM `tabStudent Master`
				WHERE IFNULL({col}, '') != '' GROUP BY {col} ORDER BY {col}""",
				as_dict=True,
			)
		return out

	where = " AND ".join(_dues_conditions(doctype))
	for key, expr in (
		("academic_year", "fd.academic_year"),
		("programme", "sm.programme_of_study"),
		("fee_component", "fd.fee_component"),
	):
		out[key] = frappe.db.sql(
			f"""SELECT {expr} AS value, COUNT(*) AS n
			FROM `tabFee Demand` fd LEFT JOIN `tabStudent Master` sm ON sm.name = fd.student
			WHERE {where} AND IFNULL({expr}, '') != ''
			GROUP BY {expr} ORDER BY {expr}""",
			as_dict=True,
		)
	return out


# ── Data ─────────────────────────────────────────────────────────────────────
DATE_BASIS = {"Due Date": "fd.due_date", "Date of creation": "fd.demand_date"}


def get_open_dues(doctype, academic_year=None, programme=None, fee_component=None, student=None,
	from_date=None, to_date=None, date_based_on="Due Date"):
	"""Open dues for a pre-filled sheet. Each filter accepts several values; empty means all.
	from_date / to_date limit the dues by Due Date or Date of creation (``date_based_on``)."""
	conditions = list(_dues_conditions(doctype))
	values = {}
	date_col = DATE_BASIS.get(date_based_on or "Due Date", "fd.due_date")
	if from_date:
		conditions.append(f"{date_col} >= %(from_date)s")
		values["from_date"] = frappe.utils.getdate(from_date)
	if to_date:
		conditions.append(f"{date_col} <= %(to_date)s")
		values["to_date"] = frappe.utils.getdate(to_date)
	for value, column, key in (
		(academic_year, "fd.academic_year", "academic_year"),
		(programme, "sm.programme_of_study", "programme"),
		(fee_component, "fd.fee_component", "fee_component"),
		(student, "fd.student", "student"),
	):
		selected = as_list(value)
		if selected:
			conditions.append(f"{column} IN %({key})s")
			values[key] = tuple(selected)

	return frappe.db.sql(
		f"""SELECT fd.name, fd.fee_component, fd.demand_date, fd.academic_year, fd.due_date,
			IFNULL(NULLIF(fd.remarks, ''), fd.description) AS remark,
			fd.original_amount, fd.penalty_amount, fd.waiver_amount, fd.net_payable,
			fd.paid_amount, fd.outstanding_amount,
			IFNULL(NULLIF(sm.registration_id, ''), sm.application_number) AS registration_id,
			IFNULL(NULLIF(sm.official_email_id, ''), sm.email) AS student_email
		FROM `tabFee Demand` fd
		LEFT JOIN `tabStudent Master` sm ON sm.name = fd.student
		WHERE {" AND ".join(conditions)}
		ORDER BY registration_id, fd.due_date, fd.name""",
		values,
		as_dict=True,
	)


def _field_options(doctype, fieldname):
	df = frappe.get_meta(doctype).get_field(fieldname)
	return [o for o in ((df.options if df else "") or "").split("\n") if o]


# ── Download ─────────────────────────────────────────────────────────────────
@frappe.whitelist()
def download_template(doctype, academic_year=None, programme=None, fee_component=None, file_type="Excel",
	from_date=None, to_date=None, date_based_on="Due Date"):
	"""The bulk-upload sheet for ``doctype`` as Excel (styled, with dropdowns) or CSV (same columns)."""
	_check(doctype)
	csv = (file_type or "").upper() == "CSV"
	if doctype == "Fee Demand":
		return _demand_template(academic_year, programme, csv)

	cfg = PREFILLED[doctype]
	dues = get_open_dues(doctype, academic_year, programme, fee_component,
		from_date=from_date, to_date=to_date, date_based_on=date_based_on)
	headers = [h for h, _w, _k in DUES_COLUMNS] + [h for h, _w, _k in cfg["inputs"]]
	rows = [[d.get(key) for _h, _w, key in DUES_COLUMNS] + [None] * len(cfg["inputs"]) for d in dues]
	for row in rows:
		for i in range(8, 14):  # amounts
			row[i] = flt(row[i])
	filename = "{0} - {1}".format(cfg["filename"], today())
	if csv:
		return _csv_response(headers, rows, filename)

	wb, ws, col = _new_sheet("Dues Details", headers, [w for _h, w, _k in DUES_COLUMNS] + [w for _h, w, _k in cfg["inputs"]])
	yellow = _fill("FFFF00")
	for h, _w, _k in cfg["inputs"]:
		ws[f"{col[h]}1"].fill = yellow
	for row in rows:
		ws.append(row)
	last = max(ws.max_row, 2)
	_format(ws, col, last, dates=("Date of creation", "Due Date"), money=[h for h, _w, _k in DUES_COLUMNS[8:]])
	for h, _w, kind in cfg["inputs"]:
		if kind == "money":
			_format(ws, col, last, money=[h])
			_decimal_validation(ws, col[h], last)
		elif kind == "date":
			_format(ws, col, last, dates=[h])
		elif kind and kind.startswith("select:"):
			_list_validation(wb, ws, col[h], last, h, _field_options(doctype, kind.split(":", 1)[1]))
	return _xlsx_response(wb, filename)


def _demand_template(academic_year, programme, csv):
	"""Blank Fee Demand creation sheet + a Students reference sheet."""
	headers = [h for h, _w, _r, _k in DEMAND_COLUMNS]
	filename = "Fee Demand Bulk Upload - {0}".format(today())
	if csv:
		return _csv_response(headers, [], filename)

	from openpyxl.comments import Comment
	from openpyxl.styles import Font

	wb, ws, col = _new_sheet("Fee Demand", headers, [w for _h, w, _r, _k in DEMAND_COLUMNS])
	last = DEMAND_ROWS + 1
	for h, _w, required, kind in DEMAND_COLUMNS:
		cell = ws[f"{col[h]}1"]
		cell.font = Font(bold=True, color="FF0000" if required and kind != "auto" else "000000")
		if kind == "auto":
			cell.font = Font(bold=True, color="808080")
			cell.comment = Comment(_("Auto-filled by the system when the demand is created. Leave blank."), "SLCM")
		elif h in ("Student Email ID", "Student Name"):
			cell.comment = Comment(_("Filled from the Student ID. If entered, the email must match that student."), "SLCM")
		elif h == "Status":
			cell.comment = Comment(_("Set by the system (Pending / Overdue …). Leave blank."), "SLCM")
		elif h == "Net Payable":
			cell.comment = Comment(_("Original − Waiver + Penalty; recalculated by the system."), "SLCM")
		elif h == "Academic Year":
			cell.comment = Comment(_("Optional — defaults to the student's academic year."), "SLCM")

	g, hh, i = col["Original Amount"], col["Waiver Amount"], col["Penalty Amount"]
	for r in range(2, last + 1):
		ws[f"{col['Net Payable']}{r}"] = f'=IF({g}{r}="","",{g}{r}-N({hh}{r})+N({i}{r}))'
	_format(ws, col, last, dates=("Demand Date", "Due Date"),
		money=("Original Amount", "Waiver Amount", "Penalty Amount", "Net Payable"))
	for h in ("Original Amount", "Waiver Amount", "Penalty Amount"):
		_decimal_validation(ws, col[h], last, allow_zero=h != "Original Amount")

	# Students reference sheet (Student ID dropdown points at it)
	students = _students(academic_year, programme)
	ref = wb.create_sheet("Students")
	ref.append(["Student ID", "Registration Id / Application Number", "Student Name", "Student Email ID", "Programme", "Academic Year"])
	for c in ref[1]:
		c.font = Font(bold=True)
	for s in students:
		ref.append([s.name, s.registration_id, s.first_name, s.email, s.programme_of_study, s.academic_year])
	for letter, width in zip("ABCDEF", (18, 30, 28, 34, 22, 14)):
		ref.column_dimensions[letter].width = width
	ref.freeze_panes = "A2"
	if students:
		_range_validation(ws, col["Student ID"], last, f"=Students!$A$2:$A${len(students) + 1}", "Student ID")
	_list_validation(wb, ws, col["Academic Year"], last, "Academic Year", frappe.get_all("Academic Year", pluck="name", order_by="name desc"))
	_list_validation(wb, ws, col["Fee Component"], last, "Fee Component", frappe.get_all("Fee Component", pluck="name", order_by="name"))
	return _xlsx_response(wb, filename)


def _students(academic_year=None, programme=None):
	conditions, values = ["1=1"], {}
	for value, column, key in ((academic_year, "academic_year", "ay"), (programme, "programme_of_study", "pr")):
		selected = as_list(value)
		if selected:
			conditions.append(f"{column} IN %({key})s")
			values[key] = tuple(selected)
	return frappe.db.sql(
		f"""SELECT name, IFNULL(NULLIF(registration_id, ''), application_number) AS registration_id, first_name,
			IFNULL(NULLIF(official_email_id, ''), email) AS email, programme_of_study, academic_year
		FROM `tabStudent Master` WHERE {" AND ".join(conditions)} ORDER BY registration_id, name""",
		values,
		as_dict=True,
	)


# ── Sheet helpers ────────────────────────────────────────────────────────────
def _fill(rgb):
	from openpyxl.styles import PatternFill

	return PatternFill("solid", fgColor=rgb)


def _new_sheet(title, headers, widths):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font
	from openpyxl.utils import get_column_letter

	wb = Workbook()
	ws = wb.active
	ws.title = title
	ws.append(headers)
	col = {}
	for idx, (header, width) in enumerate(zip(headers, widths), start=1):
		letter = get_column_letter(idx)
		col[header] = letter
		cell = ws[f"{letter}1"]
		cell.font = Font(bold=True)
		cell.alignment = Alignment(vertical="center", wrap_text=True)
		ws.column_dimensions[letter].width = width
	ws.row_dimensions[1].height = 32
	ws.freeze_panes = "A2"
	return wb, ws, col


def _format(ws, col, last, dates=(), money=()):
	for h in dates:
		for r in range(2, last + 1):
			ws[f"{col[h]}{r}"].number_format = "dd-mm-yyyy"
	for h in money:
		for r in range(2, last + 1):
			ws[f"{col[h]}{r}"].number_format = "#,##0.00"


def _decimal_validation(ws, letter, last, allow_zero=False):
	from openpyxl.worksheet.datavalidation import DataValidation

	dv = DataValidation(type="decimal", operator="greaterThanOrEqual" if allow_zero else "greaterThan", formula1="0", allow_blank=True)
	dv.error = _("Enter an amount of 0 or more.") if allow_zero else _("Enter an amount greater than zero.")
	ws.add_data_validation(dv)
	dv.add(f"{letter}2:{letter}{last + 500}")


def _range_validation(ws, letter, last, formula, title):
	from openpyxl.worksheet.datavalidation import DataValidation

	dv = DataValidation(type="list", formula1=formula, allow_blank=True)
	dv.error = _("Choose a value from the list.")
	dv.errorTitle = _("Invalid {0}").format(title)
	ws.add_data_validation(dv)
	dv.add(f"{letter}2:{letter}{last + 500}")


def _list_validation(wb, ws, letter, last, title, options):
	"""Dropdown for a column; options live on a hidden sheet so long / comma-containing values work."""
	if not options:
		return
	name = "Options"
	opt = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
	opt.sheet_state = "hidden"
	c = opt.max_column + 1 if opt.max_row > 1 or opt["A1"].value else 1
	from openpyxl.utils import get_column_letter

	ol = get_column_letter(c)
	for i, o in enumerate(options, start=1):
		opt[f"{ol}{i}"] = o
	_range_validation(ws, letter, last, f"={name}!${ol}$1:${ol}${len(options)}", title)


def _csv_response(headers, rows, filename):
	from frappe.utils.csvutils import build_csv_response

	fmt = lambda v: v.isoformat() if hasattr(v, "isoformat") else ("" if v is None else v)
	build_csv_response([headers] + [[fmt(v) for v in r] for r in rows], filename)


def _xlsx_response(wb, filename):
	buf = BytesIO()
	wb.save(buf)
	frappe.local.response.filename = f"{filename}.xlsx"
	frappe.local.response.filecontent = buf.getvalue()
	frappe.local.response.type = "binary"


# ── Data Import integration ──────────────────────────────────────────────────
@frappe.whitelist()
def data_import_download_template(doctype, export_fields=None, export_records=None, export_filters=None, file_type="CSV"):
	"""Override of Data Import's "Download Template" (hooks.py): for Fee Concession / Fee Payment /
	Fee Demand the blank template is the bulk-upload sheet. Other doctypes / export types are untouched."""
	if (doctype in PREFILLED or doctype == "Fee Demand") and export_records in (None, "", "blank_template"):
		return download_template(doctype, file_type="CSV" if (file_type or "").upper() == "CSV" else "Excel")

	from frappe.core.doctype.data_import.data_import import download_template as standard

	return standard(doctype, export_fields, export_records, export_filters, file_type)
